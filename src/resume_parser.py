from io import BytesIO
from pathlib import Path
class ResumeParseError(ValueError): pass
def extract_resume_text(filename: str,content: bytes) -> str:
 suffix=Path(filename).suffix.lower()
 if not content: raise ResumeParseError("The selected file is empty.")
 try:
  if suffix==".txt": return content.decode("utf-8-sig",errors="replace")
  if suffix==".pdf":
   from pypdf import PdfReader
   return "\n".join(p.extract_text() or "" for p in PdfReader(BytesIO(content),strict=False).pages)
  if suffix==".docx":
   from docx import Document
   d=Document(BytesIO(content));parts=[p.text for p in d.paragraphs if p.text.strip()]
   for t in d.tables: parts.extend(" | ".join(c.text for c in row.cells) for row in t.rows)
   return "\n".join(parts)
 except Exception as e: raise ResumeParseError(f"Could not read {suffix[1:].upper()} file. Check it is valid and unprotected.") from e
 raise ResumeParseError("Unsupported file type. Upload PDF, DOCX or TXT.")
